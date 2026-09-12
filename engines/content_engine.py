"""Content engine for structured text extraction with source tracking.

Provides UniversalContentEngine which extracts text from documents and returns
structured ContentBlock objects that track the source location (page, slide,
sheet, section) of each piece of content.
"""

import hashlib
import logging
import os
from dataclasses import dataclass

from text_extraction import ExtractorFactory

logger = logging.getLogger(__name__)

# Minimum OCR length below which a vision caption is attempted for embedded
# document images (short OCR usually means the image is a diagram/picture).
_EMBEDDED_IMAGE_VISION_THRESHOLD = 30


@dataclass
class ContentBlock:
    """A block of extracted text with source provenance.

    Attributes:
        text: The extracted text content.
        source_type: Type of source (page, slide, sheet, file, section, timestamp, ocr, caption, transcript, keyframe).
        source_index: Zero-based index of the source within the document.
        source_label: Human-readable label (e.g., 'Page 1', 'Slide 3', '0:30 - 1:00').
        file_path: Absolute path to the source file.
        modality: Content modality ('document', 'image', 'audio', 'video').
        timestamp_start: Start timestamp in seconds (for audio/video).
        timestamp_end: End timestamp in seconds (for audio/video).
        confidence: Extraction confidence score (0.0 - 1.0).
        metadata: Additional metadata dict.
    """

    text: str
    source_type: str
    source_index: int
    source_label: str
    file_path: str
    modality: str = "document"
    timestamp_start: float = 0.0
    timestamp_end: float = 0.0
    confidence: float = 1.0
    metadata: dict = None

    def __post_init__(self):
        if self.metadata is None:
            self.metadata = {}


class UniversalContentEngine:
    """Extracts structured content blocks from supported file types.

    Uses text_extraction.ExtractorFactory for basic extraction, then wraps
    results with source tracking metadata. For formats that support
    page/slide/sheet-level extraction, returns one ContentBlock per unit.
    """

    # Section chunk size for formats that require splitting
    _SECTION_SIZE = 500

    def extract(self, file_path: str) -> list[ContentBlock]:
        """Extract structured content blocks from a file.

        Args:
            file_path: Path to the file to extract content from.

        Returns:
            List of ContentBlock objects with source tracking.
            Returns empty list on error or unsupported file types.
        """
        file_path = os.path.abspath(file_path)
        logger.info("ContentEngine: starting extraction for '%s'", file_path)

        try:
            if not os.path.isfile(file_path):
                logger.error("ContentEngine: file not found: '%s'", file_path)
                return []

            # Compute file hash
            file_hash = self._compute_hash(file_path)
            logger.debug("ContentEngine: file hash = %s", file_hash)

            from engines.config import detect_modality_and_ext
            modality, ext = detect_modality_and_ext(file_path)

            # Route to appropriate handler
            if modality == "image":
                blocks = self._extract_image(file_path)
            elif ext in ('.mp3', '.wav', '.flac', '.ogg', '.aac',
                         '.m4a', '.wma', '.opus'):
                blocks = self._extract_audio(file_path)
            elif ext in ('.mp4', '.mkv', '.mov', '.avi', '.webm',
                         '.m4v', '.flv', '.wmv'):
                blocks = self._extract_video(file_path)
            elif ext == '.pdf':
                blocks = self._extract_pdf(file_path)
            elif ext in ('.pptx', '.ppt'):
                blocks = self._extract_pptx(file_path)
            elif ext in ('.xlsx', '.xls'):
                blocks = self._extract_xlsx(file_path)
            elif ext in ('.csv', '.tsv'):
                blocks = self._extract_csv(file_path)
            elif ext == '.docx':
                blocks = self._extract_docx(file_path)
            elif ext in ('.txt', '.text', '.log', '.md', '.markdown', '.mdown',
                         '.json', '.xml', '.xhtml', '.svg',
                         '.py', '.js', '.ts', '.java', '.cpp', '.c', '.h',
                         '.rs', '.go', '.html', '.css', '.yaml', '.yml',
                         '.toml', '.ini', '.cfg', '.conf', '.sh', '.bash',
                         '.rb', '.php', '.sql') or not ext:
                blocks = self._extract_text_chunked(file_path)
            else:
                logger.warning(
                    "ContentEngine: unsupported extension '%s' (modality: %s)", ext, modality
                )
                return []

            for block in blocks:
                if block.metadata is None:
                    block.metadata = {}
                block.metadata["file_hash"] = file_hash

            logger.info(
                "ContentEngine: finished extraction for '%s' — %d blocks",
                file_path, len(blocks)
            )
            return blocks

        except Exception as exc:
            logger.error(
                "ContentEngine: unexpected error extracting '%s': %s",
                file_path, exc, exc_info=True
            )
            return []

    def extract_full_text(self, file_path: str) -> str:
        """Extract all text from a file as a single concatenated string.

        Used for Tier 1 direct-to-LLM approach where the full file content
        is sent without chunking. Includes source labels for context.

        Args:
            file_path: Path to the file to extract text from.

        Returns:
            Complete text content of the file with source annotations.
            Returns empty string on error or unsupported file types.
        """
        blocks = self.extract(file_path)
        if not blocks:
            return ""

        parts = []
        for block in blocks:
            label = block.source_label
            text = block.text.strip()
            if not text:
                continue

            # Annotate with source label for context
            if label and label not in ("File", "Filename"):
                parts.append(f"[{label}]\n{text}")
            else:
                parts.append(text)

        return "\n\n".join(parts)

    def _compute_hash(self, file_path: str) -> str:
        """Compute SHA-256 hash of a file.

        Args:
            file_path: Path to the file.

        Returns:
            Hex digest of the SHA-256 hash.
        """
        hasher = hashlib.sha256()
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(65536), b''):
                hasher.update(chunk)
        return hasher.hexdigest()

    def _extract_pdf(self, file_path: str) -> list[ContentBlock]:
        """Extract one ContentBlock per page from a PDF file.

        Embedded raster images on each page are additionally OCR'd / captioned
        so diagrams and screenshots inside PDFs become searchable.
        """
        try:
            import fitz  # PyMuPDF
        except ImportError:
            logger.error("ContentEngine: PyMuPDF (fitz) not installed")
            return []

        blocks: list[ContentBlock] = []
        seen_xrefs: set = set()
        doc = None
        try:
            doc = fitz.open(file_path)
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text()
                if text.strip():
                    blocks.append(ContentBlock(
                        text=text,
                        source_type='page',
                        source_index=page_num,
                        source_label=f'Page {page_num + 1}',
                        file_path=file_path,
                    ))

                # Embedded images on this page (deduped across pages)
                image_blocks = self._extract_embedded_images(
                    file_path, page, page_num, 'page', seen_xrefs
                )
                blocks.extend(image_blocks)
        except Exception as exc:
            logger.error(
                "ContentEngine: error reading PDF '%s': %s",
                file_path, exc
            )
        finally:
            try:
                if doc is not None:
                    doc.close()
            except Exception:
                pass
        return blocks

    def _extract_pptx(self, file_path: str) -> list[ContentBlock]:
        """Extract one ContentBlock per slide from a PPTX file."""
        try:
            from pptx import Presentation
        except ImportError:
            logger.error("ContentEngine: python-pptx not installed")
            return []

        blocks: list[ContentBlock] = []
        try:
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides):
                texts: list[str] = []
                image_index = 0
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            para_text = paragraph.text.strip()
                            if para_text:
                                texts.append(para_text)
                    # Embedded pictures on this slide
                    image = getattr(shape, "image", None)
                    if image is not None:
                        try:
                            data = image.blob
                        except Exception:
                            data = None
                        if data:
                            blocks.extend(self._embedded_image_blocks(
                                file_path, data,
                                source_type='embedded_image',
                                source_index=slide_idx,
                                source_label=f'Slide {slide_idx + 1} · Image {image_index + 1}',
                                metadata={'slide': slide_idx, 'image_index': image_index},
                            ))
                            image_index += 1
                slide_text = '\n'.join(texts)
                if slide_text.strip():
                    blocks.append(ContentBlock(
                        text=slide_text,
                        source_type='slide',
                        source_index=slide_idx,
                        source_label=f'Slide {slide_idx + 1}',
                        file_path=file_path,
                    ))
        except Exception as exc:
            logger.error(
                "ContentEngine: error reading PPTX '%s': %s",
                file_path, exc
            )
        return blocks

    def _extract_xlsx(self, file_path: str) -> list[ContentBlock]:
        """Extract one ContentBlock per sheet from an XLSX file."""
        try:
            import openpyxl
        except ImportError:
            logger.error("ContentEngine: openpyxl not installed")
            return []

        blocks: list[ContentBlock] = []
        wb = None
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sheet_idx, sheet_name in enumerate(wb.sheetnames):
                ws = wb[sheet_name]
                rows: list[str] = []
                for row in ws.iter_rows(values_only=True):
                    cell_values = [
                        str(cell) for cell in row if cell is not None
                    ]
                    if cell_values:
                        rows.append('\t'.join(cell_values))
                sheet_text = '\n'.join(rows)
                if sheet_text.strip():
                    blocks.append(ContentBlock(
                        text=sheet_text,
                        source_type='sheet',
                        source_index=sheet_idx,
                        source_label=f'Sheet {sheet_idx + 1}',
                        file_path=file_path,
                    ))
        except Exception as exc:
            logger.error(
                "ContentEngine: error reading XLSX '%s': %s",
                file_path, exc
            )
        finally:
            try:
                if wb is not None:
                    wb.close()
            except Exception:
                pass
        return blocks

    def _extract_csv(self, file_path: str) -> list[ContentBlock]:
        """Extract entire CSV/TSV file as a single ContentBlock."""
        blocks: list[ContentBlock] = []
        try:
            extractor = ExtractorFactory.get_extractor(file_path)
            if extractor is None:
                return []
            text = extractor.extract(file_path)
            if text.strip():
                blocks.append(ContentBlock(
                    text=text,
                    source_type='file',
                    source_index=0,
                    source_label='File',
                    file_path=file_path,
                ))
        except Exception as exc:
            logger.error(
                "ContentEngine: error reading CSV '%s': %s",
                file_path, exc
            )
        return blocks

    def _extract_docx(self, file_path: str) -> list[ContentBlock]:
        """Extract ContentBlocks from DOCX, grouped into ~500 char sections."""
        try:
            import docx
        except ImportError:
            logger.error("ContentEngine: python-docx not installed")
            return []

        blocks: list[ContentBlock] = []
        try:
            doc = docx.Document(file_path)
            current_text = ''
            section_idx = 0

            # Embedded images anywhere in the document (inline + anchored)
            for part in doc.part.related_parts.values():
                content_type = getattr(part, "content_type", "") or ""
                if content_type.startswith("image/"):
                    try:
                        data = part.blob
                    except Exception:
                        data = None
                    if data:
                        blocks.extend(self._embedded_image_blocks(
                            file_path, data,
                            source_type='embedded_image',
                            source_index=len(blocks),
                            source_label=f'Image {len([b for b in blocks if b.source_type == "embedded_image"]) + 1}',
                            metadata={'image_index': len([b for b in blocks if b.source_type == "embedded_image"])},
                        ))

            for para in doc.paragraphs:
                para_text = para.text.strip()
                if not para_text:
                    continue

                if len(current_text) + len(para_text) + 1 > self._SECTION_SIZE and current_text:
                    blocks.append(ContentBlock(
                        text=current_text,
                        source_type='section',
                        source_index=section_idx,
                        source_label=f'Section {section_idx + 1}',
                        file_path=file_path,
                    ))
                    section_idx += 1
                    current_text = para_text
                else:
                    if current_text:
                        current_text += '\n' + para_text
                    else:
                        current_text = para_text

            # Flush remaining text
            if current_text.strip():
                blocks.append(ContentBlock(
                    text=current_text,
                    source_type='section',
                    source_index=section_idx,
                    source_label=f'Section {section_idx + 1}',
                    file_path=file_path,
                ))
        except Exception as exc:
            logger.error(
                "ContentEngine: error reading DOCX '%s': %s",
                file_path, exc
            )
        return blocks

    def _extract_text_chunked(self, file_path: str) -> list[ContentBlock]:
        """Extract text-based files and split into ~500 char sections."""
        blocks: list[ContentBlock] = []
        try:
            extractor = ExtractorFactory.get_extractor(file_path)
            if extractor is None:
                return []
            text = extractor.extract(file_path)
            if not text.strip():
                return []

            # Split into chunks of approximately _SECTION_SIZE characters
            chunks = self._chunk_text(text, self._SECTION_SIZE)
            for idx, chunk in enumerate(chunks):
                if chunk.strip():
                    blocks.append(ContentBlock(
                        text=chunk,
                        source_type='section',
                        source_index=idx,
                        source_label=f'Section {idx + 1}',
                        file_path=file_path,
                    ))
        except Exception as exc:
            logger.error(
                "ContentEngine: error reading text file '%s': %s",
                file_path, exc
            )
        return blocks

    def _extract_image(self, file_path: str) -> list[ContentBlock]:
        """Extract content from an image file using OCR and vision."""
        blocks: list[ContentBlock] = []
        try:
            from extractors.image_extractor import ImageExtractor
            extractor = ImageExtractor()
            blocks = extractor.extract(file_path)
        except Exception as exc:
            logger.error(
                "ContentEngine: error extracting image '%s': %s",
                file_path, exc
            )
        return blocks

    def _extract_audio(self, file_path: str) -> list[ContentBlock]:
        """Extract content from an audio file using Whisper transcription."""
        blocks: list[ContentBlock] = []
        try:
            from extractors.audio_extractor import AudioExtractor
            extractor = AudioExtractor()
            blocks = extractor.extract(file_path)
        except Exception as exc:
            logger.error(
                "ContentEngine: error extracting audio '%s': %s",
                file_path, exc
            )
        return blocks

    def _extract_video(self, file_path: str) -> list[ContentBlock]:
        """Extract content from a video file (audio transcription + keyframes)."""
        blocks: list[ContentBlock] = []
        try:
            from extractors.video_extractor import VideoExtractor
            extractor = VideoExtractor()
            blocks = extractor.extract(file_path)
        except Exception as exc:
            logger.error(
                "ContentEngine: error extracting video '%s': %s",
                file_path, exc
            )
        return blocks

    # ------------------------------------------------------------------ #
    # Embedded image extraction (PDF / PPTX / DOCX)
    # ------------------------------------------------------------------ #
    def _extract_embedded_images(
        self, file_path: str, page, page_num: int, source_type: str,
        seen_xrefs: set = None,
    ) -> list[ContentBlock]:
        """Extract and OCR/caption raster images embedded in a PDF page.

        Args:
            seen_xrefs: Optional document-level set of xrefs already OCR'd.
                Shared image resources are processed only once per document.
        """
        blocks: list[ContentBlock] = []
        if seen_xrefs is None:
            seen_xrefs = set()
        try:
            images = page.get_images(full=True)
            for img_idx, img_info in enumerate(images):
                xref = img_info[0]
                # Same shared image resource can appear on many pages; OCR it
                # only once per document to avoid duplicate indexing.
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                data = None
                try:
                    doc = page.parent if hasattr(page, "parent") else None
                    if doc is None:
                        continue
                    extracted = doc.extract_image(xref)
                    if extracted and extracted.get("image"):
                        data = extracted["image"]
                except Exception:
                    data = None

                if not data:
                    continue

                label = f'Page {page_num + 1} · Image {img_idx + 1}'
                blocks.extend(self._embedded_image_blocks(
                    file_path, data,
                    source_type='embedded_image',
                    source_index=page_num,
                    source_label=label,
                    metadata={'page': page_num, 'image_index': img_idx},
                ))
        except Exception as exc:
            logger.debug("Embedded image extraction failed for PDF %s: %s", file_path, exc)
        return blocks

    def _embedded_image_blocks(
        self, file_path: str, data: bytes,
        source_type: str, source_index: int, source_label: str, metadata: dict,
    ) -> list[ContentBlock]:
        """OCR + (if needed) vision-caption raw image bytes into ContentBlocks.

        OCR and vision run independently: if OCR fails, the vision caption is
        still attempted (diagrams often produce no OCR but a useful caption).
        """
        blocks: list[ContentBlock] = []

        # OCR pass (independent — a tesseract failure must not block captioning)
        ocr_text = ""
        try:
            from extractors.image_extractor import ocr_image_bytes
            ocr_text = ocr_image_bytes(data) or ""
        except Exception as exc:
            logger.debug("Embedded image OCR failed: %s", exc)

        if ocr_text.strip():
            blocks.append(ContentBlock(
                text=ocr_text,
                source_type='ocr',
                source_index=source_index,
                source_label=f'{source_label} · OCR',
                file_path=file_path,
                modality='document',
                metadata={**metadata, 'embedded': True},
                confidence=0.8,
            ))

        # Vision caption when OCR found little text (likely a diagram)
        if len(ocr_text.strip()) < _EMBEDDED_IMAGE_VISION_THRESHOLD:
            try:
                from extractors.image_extractor import caption_image_bytes
                caption = caption_image_bytes(data) or ""
            except Exception as exc:
                logger.debug("Embedded image vision caption failed: %s", exc)
                caption = ""
            if caption.strip():
                blocks.append(ContentBlock(
                    text=caption,
                    source_type='caption',
                    source_index=source_index,
                    source_label=f'{source_label} · Caption',
                    file_path=file_path,
                    modality='document',
                    metadata={**metadata, 'embedded': True},
                    confidence=0.7,
                ))
        return blocks

    @staticmethod
    def _chunk_text(text: str, chunk_size: int) -> list[str]:
        """Split text into chunks of approximately chunk_size characters.

        Attempts to break on newline or space boundaries.

        Args:
            text: The text to split.
            chunk_size: Target size for each chunk.

        Returns:
            List of text chunks.
        """
        chunks: list[str] = []
        start = 0
        text_len = len(text)

        while start < text_len:
            end = start + chunk_size

            if end >= text_len:
                chunks.append(text[start:])
                break

            # Try to break at a newline within the last 20% of the chunk
            search_start = end - chunk_size // 5
            newline_pos = text.rfind('\n', search_start, end)
            if newline_pos > start:
                end = newline_pos + 1
            else:
                # Try to break at a space
                space_pos = text.rfind(' ', search_start, end)
                if space_pos > start:
                    end = space_pos + 1

            chunks.append(text[start:end])
            start = end

        return chunks
