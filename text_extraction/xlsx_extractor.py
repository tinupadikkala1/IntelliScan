"""XLSX file extractor using openpyxl."""

import logging

from openpyxl import load_workbook

from .base_extractor import BaseExtractor

logger = logging.getLogger(__name__)


class XlsxExtractor(BaseExtractor):
    """Extractor for XLSX (Microsoft Excel) files.

    Uses openpyxl to read all sheets and all rows.
    """

    def extract(self, file_path: str) -> str:
        """Extract text from an XLSX file.

        Args:
            file_path: Path to the XLSX file.

        Returns:
            Extracted text content as a string. Returns empty string on failure.
        """
        logger.info("Starting XLSX extraction: %s", file_path)
        try:
            if not self.validate_file(file_path):
                return ""

            wb = load_workbook(file_path, read_only=True, data_only=True)
            text_parts: list[str] = []

            try:
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    text_parts.append(f"--- Sheet: {sheet_name} ---")

                    for row in sheet.iter_rows(values_only=True):
                        row_values = [
                            str(cell) if cell is not None else ""
                            for cell in row
                        ]
                        row_text = "\t".join(row_values)
                        if row_text.strip():
                            text_parts.append(row_text)
            finally:
                wb.close()

            result = "\n".join(text_parts)
            logger.info("Finished XLSX extraction: %s", file_path)
            return result

        except Exception as e:
            logger.error("Error extracting XLSX from %s: %s", file_path, str(e))
            return ""
