"""PPTX file extractor using python-pptx."""

import logging

from pptx import Presentation

from .base_extractor import BaseExtractor

logger = logging.getLogger(__name__)


class PptxExtractor(BaseExtractor):
    """Extractor for PPTX (Microsoft PowerPoint) files.

    Uses python-pptx to extract text from all slides and shapes.
    """

    def extract(self, file_path: str) -> str:
        """Extract text from a PPTX file.

        Args:
            file_path: Path to the PPTX file.

        Returns:
            Extracted text content as a string. Returns empty string on failure.
        """
        logger.info("Starting PPTX extraction: %s", file_path)
        try:
            if not self.validate_file(file_path):
                return ""

            prs = Presentation(file_path)
            text_parts: list[str] = []

            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_texts: list[str] = []

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            para_text = paragraph.text.strip()
                            if para_text:
                                slide_texts.append(para_text)

                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            row_text = "\t".join(
                                cell.text.strip() for cell in row.cells
                            )
                            if row_text.strip():
                                slide_texts.append(row_text)

                if slide_texts:
                    text_parts.append(f"--- Slide {slide_num} ---")
                    text_parts.extend(slide_texts)

            result = "\n".join(text_parts)
            logger.info("Finished PPTX extraction: %s", file_path)
            return result

        except Exception as e:
            logger.error("Error extracting PPTX from %s: %s", file_path, str(e))
            return ""
