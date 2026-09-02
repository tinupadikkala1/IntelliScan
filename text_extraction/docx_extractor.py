"""DOCX file extractor using python-docx."""

import logging

from docx import Document

from .base_extractor import BaseExtractor

logger = logging.getLogger(__name__)


class DocxExtractor(BaseExtractor):
    """Extractor for DOCX (Microsoft Word) files.

    Uses python-docx to extract text from paragraphs and tables.
    """

    def extract(self, file_path: str) -> str:
        """Extract text from a DOCX file.

        Args:
            file_path: Path to the DOCX file.

        Returns:
            Extracted text content as a string. Returns empty string on failure.
        """
        logger.info("Starting DOCX extraction: %s", file_path)
        try:
            if not self.validate_file(file_path):
                return ""

            doc = Document(file_path)
            text_parts: list[str] = []

            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)

            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = "\t".join(
                        cell.text.strip() for cell in row.cells
                    )
                    if row_text.strip():
                        text_parts.append(row_text)

            result = "\n".join(text_parts)
            logger.info("Finished DOCX extraction: %s", file_path)
            return result

        except Exception as e:
            logger.error("Error extracting DOCX from %s: %s", file_path, str(e))
            return ""
